"""
Excel 导出服务
"""
import os
import logging
from datetime import datetime
from typing import List
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from .invoice import Invoice, LineItem


def create_new_excel(filename: str, sheet_name: str = '发票明细') -> Workbook:
    """
    创建新的 Excel 文件并写入表头
    
    Args:
        filename: Excel 文件路径
        sheet_name: 工作表名称
        
    Returns:
        Workbook 对象
    """
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    # 定义表头
    headers = [
        '发票号码', '开票日期', '购买方', '项目名称', '规格', '单位',
        '数量', '单价', '金额', '税率', '税额', '发票文件名',
        '录入时间', '购买方统一社会信用代码/纳税人识别号', '发票类型'
    ]

    # 写入表头
    for col_idx, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col_idx, value=header)

    # 保存文件
    wb.save(filename)
    return wb


def find_next_row(ws: Worksheet) -> int:
    """
    找到当前工作表中已使用的最后一行的下一行行号
    
    Args:
        ws: 工作表对象
        
    Returns:
        下一行行号（从1开始）
    """
    # 查找最后一行（基于 A 列）
    last_row = 0
    for row in ws.iter_rows(min_row=1, max_col=1):
        if row[0].value:
            last_row = row[0].row

    # 如果文件为空，last_row 为 0
    # 如果文件只有表头，last_row 为 1
    # 我们始终从 last_row + 1 开始写入
    return last_row + 1


def batch_export_to_excel(invoice_list: List[Invoice], filename: str) -> None:
    """
    批量将发票列表导出到 Excel 文件
    
    Args:
        invoice_list: 发票列表
        filename: Excel 文件路径
    """
    sheet_name = '发票明细'

    # 1. 打开现有文件或创建新文件
    if os.path.exists(filename):
        try:
            wb = load_workbook(filename)
            if sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
            else:
                ws = wb.create_sheet(sheet_name)
        except Exception as e:
            logging.warning(f'打开 Excel 文件失败，将创建新文件: {e}')
            wb = create_new_excel(filename, sheet_name)
            ws = wb[sheet_name]
    else:
        logging.info('文件不存在，正在创建新文件并写入数据...')
        wb = create_new_excel(filename, sheet_name)
        ws = wb[sheet_name]

    # 2. 确定写入起点（最后一行 + 1）
    start_row = find_next_row(ws)

    # 3. 写入数据行
    current_row = start_row
    handle_time = datetime.now().strftime('%Y-%m-%d %H:%M:%S')

    for invoice in invoice_list:
        for item in invoice.items:
            # 写入数据到单元格
            # A 列: 发票号码
            ws.cell(row=current_row, column=1, value=invoice.code)
            # B 列: 开票日期
            ws.cell(row=current_row, column=2, value=invoice.date)
            # C 列: 购买方
            ws.cell(row=current_row, column=3, value=invoice.buyer)
            # D 列: 项目名称
            ws.cell(row=current_row, column=4, value=item.item_name)
            # E 列: 规格
            ws.cell(row=current_row, column=5, value=item.spec)
            # F 列: 单位
            ws.cell(row=current_row, column=6, value=item.unit)
            # G 列: 数量
            ws.cell(row=current_row, column=7, value=item.quantity)
            # H 列: 单价
            ws.cell(row=current_row, column=8, value=item.price)
            # I 列: 金额
            ws.cell(row=current_row, column=9, value=item.amount)
            # J 列: 税率
            ws.cell(row=current_row, column=10, value=item.tax_rate)
            # K 列: 税额
            ws.cell(row=current_row, column=11, value=item.tax_amount)
            # L 列: 发票文件名
            ws.cell(row=current_row, column=12, value=invoice.name)
            # M 列: 录入时间
            ws.cell(row=current_row, column=13, value=handle_time)
            # N 列: 购买方统一社会信用代码/纳税人识别号
            ws.cell(row=current_row, column=14, value=invoice.buyer_tax_id)
            # O 列: 发票类型
            ws.cell(row=current_row, column=15, value=invoice.invoice_type)

            current_row += 1

    # 4. 保存文件到磁盘
    try:
        wb.save(filename)
        logging.info(f'数据已成功导出到: {filename}')
    except Exception as e:
        raise Exception(f'保存 Excel 文件失败: {e}')


